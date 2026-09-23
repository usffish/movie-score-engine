"""
Tests for workbook-aware manual entry and smart-update re-checks.

Covers:
- --manual only asks about scores blank in the workbook; Reviews only
  alongside a Metascore typed in the same prompt
- failed movies: only blank fields are asked; complete rows aren't asked
- the Manual column records typed scores; a scraped score replaces a typed
  one and clears it
- typed/existing scores count in the composite
- smart-update always re-checks rows with a blank or typed score
- Table1 grows to cover added columns (header names + autoFilter)
"""

import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

from excel import (
    ensure_headers,
    extend_table_to_stability_cols,
    get_header_map,
    read_manual_fields,
    should_update,
    write_manual_fields,
)
from manual import apply_manual_entry, prompt_failed_movie, prompt_missing_scores
from scoring import RawScores


def scores(title, metascore=None, imdb=None, reviews=0, lb=None, year=None):
    return RawScores(title, metascore, imdb, reviews, lb, year=year)


# ---------------------------------------------------------------------------
# Prompts
# ---------------------------------------------------------------------------

class TestPrompts(unittest.TestCase):

    def ask(self, fn, *args, inputs, **kwargs):
        prompts = []

        def fake_input(prompt):
            prompts.append(prompt.strip())
            return inputs.pop(0) if inputs else ""

        with patch("builtins.input", side_effect=fake_input), patch("builtins.print"):
            result = fn(*args, **kwargs)
        return result, prompts

    def test_reviews_not_asked_when_metascore_already_known(self):
        # Metascore came from OMDb, no Metacritic page: nothing to ask about reviews.
        raw = scores("A", metascore=70, imdb=7.0, reviews=0, lb=None)
        result, prompts = self.ask(prompt_missing_scores, raw, inputs=["3.5"])
        self.assertEqual(prompts, ["Letterboxd rating (0.0-5.0):"])
        self.assertEqual(result.letterboxd_rating, 3.5)

    def test_reviews_asked_after_typed_metascore(self):
        raw = scores("A", imdb=7.0, lb=3.5)
        result, prompts = self.ask(prompt_missing_scores, raw, inputs=["74", "12"])
        self.assertEqual(prompts, ["Metascore (0-100):", "Critic review count (0+):"])
        self.assertEqual((result.metascore, result.review_count), (74, 12))

    def test_reviews_not_asked_when_metascore_skipped(self):
        raw = scores("A", imdb=7.0, lb=3.5)
        _, prompts = self.ask(prompt_missing_scores, raw, inputs=[""])
        self.assertEqual(prompts, ["Metascore (0-100):"])

    def test_movie_missing_only_reviews_is_not_prompted(self):
        raw = scores("A", metascore=70, imdb=7.0, reviews=0, lb=3.5)
        with patch("builtins.input", side_effect=AssertionError("should not prompt")):
            updated, _, _ = apply_manual_entry([raw], [], manual=True)
        self.assertEqual(updated, [raw])

    def test_failed_movie_asks_only_blank_fields(self):
        existing = scores("X", metascore=60, reviews=8, imdb=6.5)
        result, prompts = self.ask(prompt_failed_movie, "X", existing=existing, inputs=["3.1"])
        self.assertEqual(prompts, ["Letterboxd rating (0.0-5.0):"])
        self.assertEqual(result, scores("X", metascore=60, imdb=6.5, reviews=8, lb=3.1))

    def test_complete_failed_movie_is_not_prompted(self):
        existing = {"X": scores("X", metascore=60, imdb=6.5, reviews=8, lb=3.1)}
        with patch("builtins.input", side_effect=AssertionError("should not prompt")):
            updated, failed, _ = apply_manual_entry([], ["X"], manual=True, existing=existing)
        self.assertEqual(failed, ["X"])
        self.assertEqual(updated, [])


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------

def _sheet(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    return wb, ws


class TestManualColumn(unittest.TestCase):

    def test_round_trip_in_column_order(self):
        _, ws = _sheet([["Movies", "Manual"], ["A", None]])
        hm = get_header_map(ws)
        write_manual_fields(ws, 2, hm, {"IMDB", "Metacritic"})
        self.assertEqual(ws.cell(2, 2).value, "Metacritic, IMDB")
        self.assertEqual(read_manual_fields(ws, 2, hm), {"Metacritic", "IMDB"})
        write_manual_fields(ws, 2, hm, set())
        self.assertIsNone(ws.cell(2, 2).value)

    def test_unknown_names_ignored(self):
        _, ws = _sheet([["Movies", "Manual"], ["A", "IMDB, nonsense"]])
        self.assertEqual(read_manual_fields(ws, 2, get_header_map(ws)), {"IMDB"})


class TestSmartUpdateRechecks(unittest.TestCase):

    HEADERS = ["Movies", "Metacritic", "Letterboxd", "IMDB", "TRUE", "LastUpdated", "StableWeeks", "Manual"]
    TODAY = date(2026, 9, 23)

    def _row(self, *values):
        _, ws = _sheet([self.HEADERS, list(values)])
        return ws, get_header_map(ws)

    def test_stable_complete_movie_is_skipped(self):
        ws, hm = self._row("A", 70, 3.5, 7.0, 0.6, "2026-09-20", 4, None)
        self.assertFalse(should_update(ws, 2, hm, self.TODAY))

    def test_blank_score_rechecked_even_after_previous_runs(self):
        ws, hm = self._row("A", None, 3.5, 7.0, 0.6, "2026-09-20", 4, None)
        self.assertTrue(should_update(ws, 2, hm, self.TODAY))

    def test_typed_score_rechecked(self):
        ws, hm = self._row("A", 70, 3.5, 7.0, 0.6, "2026-09-20", 4, "IMDB")
        self.assertTrue(should_update(ws, 2, hm, self.TODAY))


class TestTableGrowth(unittest.TestCase):

    def test_table_covers_appended_columns(self):
        wb, ws = _sheet([["Movies", "Metacritic"], ["A", 70], ["B", 80]])
        table = Table(displayName="Table1", ref="A1:B3")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9")
        table.autoFilter = AutoFilter(ref="A1:B3")
        table.tableColumns = [TableColumn(id=1, name="Movies"), TableColumn(id=2, name="Metacritic")]
        ws.add_table(table)

        hm = ensure_headers(ws, get_header_map(ws))  # appends Year ... Manual
        extend_table_to_stability_cols(ws)

        table = ws.tables["Table1"]
        last = hm["Manual"]
        self.assertEqual(table.ref, f"A1:{openpyxl.utils.get_column_letter(last)}3")
        self.assertEqual(table.autoFilter.ref, table.ref)
        names = [c.name for c in table.tableColumns]
        self.assertEqual(names, [ws.cell(1, c).value for c in range(1, last + 1)])
        ids = [c.id for c in table.tableColumns]
        self.assertEqual(len(ids), len(set(ids)))


# ---------------------------------------------------------------------------
# Several runs of the real workflow
# ---------------------------------------------------------------------------

class TestManualAcrossRuns(unittest.TestCase):
    """
    Run 1: IMDB missing -> typed in, flagged Manual, counted in TRUE.
    Run 2 on that output: still missing -> not asked again, value kept.
    Run 3: the source now has it -> replaces the typed value, flag cleared.
    """

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.book = self.dir / "Movies.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Movies", "Year"])
        ws.append(["Alpha", 2026])
        ws.append(["Beta", 2026])
        wb.save(self.book)

    def tearDown(self):
        self._tmp.cleanup()

    def _run(self, imdb_alpha, inputs):
        from update_scores import update_workbook

        def fetch_all(movies, api_key, delay=0.0, verbose=False, resolver=None,
                      rate_limiter=None, years=None):
            data = {
                "Alpha": scores("Alpha", 70, imdb_alpha, 20, 3.5, year=2026),
                "Beta": scores("Beta", 80, 8.0, 30, 4.0, year=2026),
            }
            return [data[m] for m in movies], []

        prompts = []

        def fake_input(prompt):
            prompts.append(prompt.strip())
            if not inputs:
                raise EOFError
            return inputs.pop(0)

        with patch("update_scores.fetch_all", side_effect=fetch_all), \
             patch("builtins.input", side_effect=fake_input), \
             patch("builtins.print"):
            update_workbook(self.book, self.book, api_key="k", delay=0.0, manual=True)

        ws = openpyxl.load_workbook(self.book).active
        hm = get_header_map(ws)
        row = {ws.cell(r, hm["Movies"]).value: r for r in range(2, ws.max_row + 1)}["Alpha"]
        cell = lambda col: ws.cell(row, hm[col]).value
        return prompts, cell

    def test_workflow(self):
        prompts, cell = self._run(imdb_alpha=None, inputs=["7.0"])
        self.assertEqual(prompts, ["IMDB rating (0.0-10.0):"])
        self.assertEqual(cell("IMDB"), 7.0)
        self.assertEqual(cell("Manual"), "IMDB")
        self.assertIsNotNone(cell("st.IMDB"))  # typed score counts in the composite

        prompts, cell = self._run(imdb_alpha=None, inputs=[])
        self.assertEqual(prompts, [])            # not asked again
        self.assertEqual(cell("IMDB"), 7.0)
        self.assertEqual(cell("Manual"), "IMDB")

        prompts, cell = self._run(imdb_alpha=7.4, inputs=[])
        self.assertEqual(prompts, [])
        self.assertEqual(cell("IMDB"), 7.4)      # real score replaces the typed one
        self.assertIsNone(cell("Manual"))


if __name__ == "__main__":
    unittest.main()
