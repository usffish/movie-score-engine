"""
Tests for release-year disambiguation.

Covers:
- Letterboxd / Metacritic try the year-suffixed slug before the plain slug
- Letterboxd / Metacritic skip a page whose release year doesn't match
- read_years parses the optional Year column
- resolve_year: user year wins, agreeing sources fill it, disagreement = unknown
- --manual asks for unknown years, re-fetches with them, and stops on EOF
- the matched year is written to the output Year column
"""

import unittest
from unittest.mock import MagicMock, patch

import openpyxl

from excel import get_header_map, read_years
from scraper.letterboxd_scraper import _candidate_slugs, get_letterboxd_data


def _page(title_with_year: str, rating: float) -> MagicMock:
    resp = MagicMock()
    resp.status_code = 200
    resp.text = f"""
    <html><head>
    <meta property="og:title" content="{title_with_year}">
    <meta itemprop="ratingValue" content="{rating}">
    </head><body></body></html>
    """
    return resp


def _not_found() -> MagicMock:
    resp = MagicMock()
    resp.status_code = 404
    resp.text = ""
    return resp


class TestCandidateSlugs(unittest.TestCase):

    def test_year_slug_comes_first(self):
        self.assertEqual(
            _candidate_slugs("Parasite", 2019)[:2],
            ["parasite-2019", "parasite"],
        )

    def test_no_year_starts_with_plain_slug(self):
        self.assertEqual(_candidate_slugs("Parasite")[0], "parasite")


class TestLetterboxdYear(unittest.TestCase):

    def _fake_get(self, pages):
        def get(url, **kwargs):
            return pages.get(url, _not_found())
        return get

    def test_year_suffixed_page_is_used(self):
        pages = {
            "https://letterboxd.com/film/parasite-2019/": _page("Parasite (2019)", 4.5),
            "https://letterboxd.com/film/parasite/": _page("Parasite (1982)", 2.4),
        }
        with patch("scraper.letterboxd_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_letterboxd_data("Parasite", year=2019)
        self.assertEqual(result["rating"], 4.5)
        self.assertEqual(result["url"], "https://letterboxd.com/film/parasite-2019/")

    def test_wrong_year_plain_slug_is_skipped(self):
        pages = {
            "https://letterboxd.com/film/parasite/": _page("Parasite (1982)", 2.4),
            "https://letterboxd.com/film/parasite-1/": _page("Parasite (2019)", 4.5),
        }
        with patch("scraper.letterboxd_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_letterboxd_data("Parasite", year=2019)
        self.assertEqual(result["rating"], 4.5)

    def test_plain_slug_used_when_year_matches(self):
        pages = {
            "https://letterboxd.com/film/the-godfather/": _page("The Godfather (1972)", 4.51),
        }
        with patch("scraper.letterboxd_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_letterboxd_data("The Godfather", year=1972)
        self.assertEqual(result["rating"], 4.51)

    def test_no_year_keeps_old_behaviour(self):
        pages = {
            "https://letterboxd.com/film/parasite/": _page("Parasite (1982)", 2.4),
        }
        with patch("scraper.letterboxd_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_letterboxd_data("Parasite")
        self.assertEqual(result["rating"], 2.4)


def _mc_page(date_published: str, score: int, reviews: int) -> MagicMock:
    resp = MagicMock()
    resp.status_code = 200
    resp.text = f"""
    <html><head><script type="application/ld+json">
    {{"@type": "Movie", "datePublished": "{date_published}",
      "aggregateRating": {{"ratingValue": {score}, "reviewCount": {reviews}}}}}
    </script></head><body></body></html>
    """
    return resp


class TestMetacriticYear(unittest.TestCase):

    def _fake_get(self, pages):
        def get(url, **kwargs):
            return pages.get(url, _not_found())
        return get

    def _get(self, pages, title, year=None):
        from scraper.metacritic_scraper import get_metacritic_data
        with patch("scraper.metacritic_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            return get_metacritic_data(title, year=year)

    def test_year_suffixed_page_is_used(self):
        pages = {
            "https://www.metacritic.com/movie/buddy-2026/": _mc_page("2026-08-28", 67, 22),
            "https://www.metacritic.com/movie/buddy/": _mc_page("2019-03-20", 76, 4),
        }
        self.assertEqual(self._get(pages, "Buddy", 2026), {"review_count": 22, "metascore": 67, "year": 2026})

    def test_wrong_year_plain_slug_is_rejected(self):
        pages = {
            "https://www.metacritic.com/movie/buddy/": _mc_page("2019-03-20", 76, 4),
        }
        self.assertEqual(self._get(pages, "Buddy", 2026), {"review_count": 0, "metascore": None, "year": None})

    def test_plain_slug_used_when_year_matches(self):
        pages = {
            "https://www.metacritic.com/movie/godfather/": _mc_page("1972-03-24", 100, 16),
        }
        self.assertEqual(self._get(pages, "The Godfather", 1972), {"review_count": 16, "metascore": 100, "year": 1972})

    def test_no_year_keeps_old_behaviour(self):
        pages = {
            "https://www.metacritic.com/movie/buddy/": _mc_page("2019-03-20", 76, 4),
        }
        self.assertEqual(self._get(pages, "Buddy"), {"review_count": 4, "metascore": 76, "year": 2019})


class TestFestivalVsReleaseYear(unittest.TestCase):
    """
    Without Blood premiered in 2024 (IMDb, Letterboxd) but Metacritic dates
    it by its 2026 theatrical release.  A 2-year gap is the same film.
    """

    def _fake_get(self, pages):
        def get(url, **kwargs):
            return pages.get(url, _not_found())
        return get

    def test_metacritic_accepts_page_two_years_off(self):
        from scraper.metacritic_scraper import get_metacritic_data
        pages = {"https://www.metacritic.com/movie/without-blood/": _mc_page("2026-01-30", 41, 10)}
        with patch("scraper.metacritic_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_metacritic_data("Without Blood", year=2024)
        self.assertEqual(result, {"review_count": 10, "metascore": 41, "year": 2026})

    def test_letterboxd_accepts_page_two_years_off(self):
        pages = {"https://letterboxd.com/film/without-blood/": _page("Without Blood (2024)", 3.02)}
        with patch("scraper.letterboxd_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_letterboxd_data("Without Blood", year=2026)
        self.assertEqual(result["rating"], 3.02)

    def test_three_years_off_is_still_a_different_film(self):
        from scraper.metacritic_scraper import get_metacritic_data
        pages = {"https://www.metacritic.com/movie/without-blood/": _mc_page("2027-01-30", 41, 10)}
        with patch("scraper.metacritic_scraper.SESSION.get", side_effect=self._fake_get(pages)):
            result = get_metacritic_data("Without Blood", year=2024)
        self.assertEqual(result["review_count"], 0)

    def test_resolve_year_accepts_two_year_spread(self):
        from scoring import resolve_year
        self.assertEqual(
            resolve_year(None, {"Metacritic": 2026, "Letterboxd": 2024, "OMDb": 2024}), 2024
        )
        self.assertIsNone(
            resolve_year(None, {"Metacritic": 2026, "Letterboxd": 2022, "OMDb": 2026})
        )


class TestOmdbYearRetry(unittest.TestCase):

    def _omdb(self, data):
        resp = MagicMock()
        resp.status_code = 200
        resp.json.return_value = data
        return resp

    def _fake_get(self, by_year, no_year):
        def get(url, params=None, **kwargs):
            if "y" in params:
                return self._omdb(by_year.get(params["y"], {"Response": "False", "Error": "Movie not found!"}))
            return self._omdb(no_year)
        return get

    def test_retries_without_year_and_accepts_close_year(self):
        from scraper.omdb_client import get_omdb_data
        film = {"Response": "True", "Year": "2024", "imdbID": "tt18398986",
                "imdbRating": "6.1", "Metascore": "N/A"}
        with patch("scraper.omdb_client.SESSION.get", side_effect=self._fake_get({}, film)) as mock_get:
            result = get_omdb_data("Without Blood", "k", year=2026)
        self.assertEqual(mock_get.call_count, 2)
        self.assertEqual(result["imdb_id"], "tt18398986")
        self.assertEqual(result["year"], 2024)

    def test_retry_rejects_a_different_film(self):
        from scraper.omdb_client import get_omdb_data
        other = {"Response": "True", "Year": "2019", "imdbID": "tt0000001", "imdbRating": "7.0"}
        with patch("scraper.omdb_client.SESSION.get", side_effect=self._fake_get({}, other)):
            result = get_omdb_data("Buddy", "k", year=2026)
        self.assertIsNone(result["imdb_id"])

    def test_no_retry_when_year_search_succeeds(self):
        from scraper.omdb_client import get_omdb_data
        film = {"Response": "True", "Year": "2026", "imdbID": "tt37281055", "imdbRating": "7.1"}
        with patch("scraper.omdb_client.SESSION.get", side_effect=self._fake_get({2026: film}, {})) as mock_get:
            result = get_omdb_data("Buddy", "k", year=2026)
        self.assertEqual(mock_get.call_count, 1)
        self.assertEqual(result["imdb_id"], "tt37281055")


class TestReadYears(unittest.TestCase):

    def _sheet(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        return ws

    def test_no_year_column_returns_empty(self):
        ws = self._sheet([["Movies"], ["Parasite"]])
        self.assertEqual(read_years(ws, get_header_map(ws), [(2, "Parasite")]), {})

    def test_parses_ints_strings_and_skips_blanks(self):
        ws = self._sheet([
            ["Movies", "Year"],
            ["Parasite", 2019],
            ["Boogie Nights", "1997"],
            ["Unknown", None],
            ["Garbage", "n/a"],
        ])
        rows = [(2, "Parasite"), (3, "Boogie Nights"), (4, "Unknown"), (5, "Garbage")]
        self.assertEqual(
            read_years(ws, get_header_map(ws), rows),
            {"Parasite": 2019, "Boogie Nights": 1997},
        )


class TestResolveYear(unittest.TestCase):

    def test_given_year_always_wins(self):
        from scoring import resolve_year
        self.assertEqual(resolve_year(2026, {"Metacritic": 2019, "OMDb": 2019}), 2026)

    def test_agreeing_sources_prefer_omdb(self):
        from scoring import resolve_year
        self.assertEqual(
            resolve_year(None, {"Metacritic": 2025, "Letterboxd": 2025, "OMDb": 2026}), 2026
        )

    def test_disagreeing_sources_are_unknown(self):
        from scoring import resolve_year
        self.assertIsNone(
            resolve_year(None, {"Metacritic": 2019, "Letterboxd": 2019, "OMDb": 2026})
        )

    def test_missing_sources_are_ignored(self):
        from scoring import resolve_year
        self.assertEqual(resolve_year(None, {"Metacritic": None, "Letterboxd": 2002}), 2002)
        self.assertIsNone(resolve_year(None, {"Metacritic": None}))


class TestPromptUnknownYears(unittest.TestCase):

    def _raw(self, title, year=None, source_years=None):
        from scoring import RawScores
        return RawScores(title, 70, 7.0, 10, 3.5, year=year, source_years=source_years or {})

    def test_asks_only_for_unknown_years(self):
        from manual import prompt_unknown_years
        raws = [
            self._raw("Buddy", None, {"Metacritic": 2019, "OMDb": 2026}),
            self._raw("Hope", None, {"Metacritic": 2021, "OMDb": 2013}),
            self._raw("Known", 2020),
        ]
        with patch("builtins.input", side_effect=["2026", ""]) as mock_input, \
             patch("builtins.print"):
            years, interrupted = prompt_unknown_years(raws)
        self.assertEqual(years, {"Buddy": 2026})
        self.assertFalse(interrupted)
        self.assertEqual(mock_input.call_count, 2)

    def test_end_of_input_stops_and_keeps_entries(self):
        from manual import prompt_unknown_years
        raws = [self._raw("Buddy"), self._raw("Hope")]
        with patch("builtins.input", side_effect=["2026", EOFError]), patch("builtins.print"):
            years, interrupted = prompt_unknown_years(raws)
        self.assertEqual(years, {"Buddy": 2026})
        self.assertTrue(interrupted)

    def test_score_prompt_end_of_input_keeps_year(self):
        from manual import prompt_missing_scores, ManualEntryInterrupted
        raw = self._raw("Buddy", 2026, {"OMDb": 2026})
        raw.imdb_rating = None
        with patch("builtins.input", side_effect=EOFError), patch("builtins.print"):
            with self.assertRaises(ManualEntryInterrupted) as ctx:
                prompt_missing_scores(raw)
        self.assertEqual(ctx.exception.partial.year, 2026)


class TestWorkbookYear(unittest.TestCase):

    def setUp(self):
        import tempfile
        from pathlib import Path
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.input = self.tmp / "in.xlsx"
        self.output = self.tmp / "out.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Movies"])
        ws.append(["Buddy"])
        ws.append(["Resident Evil"])
        wb.save(self.input)

    def tearDown(self):
        self._tmp.cleanup()

    def _output_years(self):
        ws = openpyxl.load_workbook(self.output).active
        hm = get_header_map(ws)
        return {ws.cell(r, hm["Movies"]).value: ws.cell(r, hm["Year"]).value
                for r in range(2, ws.max_row + 1)}

    def _fake_fetch_all(self, calls):
        from scoring import RawScores

        def fetch_all(movies, api_key, delay=0.0, verbose=False, resolver=None,
                      rate_limiter=None, years=None):
            years = years or {}
            calls.append((list(movies), dict(years)))
            out = []
            for t in movies:
                if t == "Buddy" and years.get(t) == 2026:
                    out.append(RawScores(t, 67, 7.1, 22, 3.09, year=2026,
                                         source_years={"Metacritic": 2026, "OMDb": 2026}))
                elif t == "Buddy":
                    out.append(RawScores(t, 76, 7.1, 4, 2.54, year=None,
                                         source_years={"Metacritic": 2019, "OMDb": 2026}))
                else:
                    out.append(RawScores(t, 35, 6.6, 24, 2.94, year=2002,
                                         source_years={"Metacritic": 2002, "OMDb": 2002}))
            return out, []
        return fetch_all

    def test_matched_year_written_unknown_left_blank(self):
        from update_scores import update_workbook
        calls = []
        with patch("update_scores.fetch_all", side_effect=self._fake_fetch_all(calls)):
            update_workbook(self.input, self.output, api_key="k", delay=0.0)
        self.assertEqual(self._output_years(), {"Buddy": None, "Resident Evil": 2002})

    def test_manual_year_refetches_with_that_year(self):
        from update_scores import update_workbook
        calls = []
        # Year prompt for Buddy only (Resident Evil's year is known); then EOF
        # ends the score prompts.
        with patch("update_scores.fetch_all", side_effect=self._fake_fetch_all(calls)), \
             patch("builtins.input", side_effect=["2026", EOFError]), \
             patch("builtins.print"):
            update_workbook(self.input, self.output, api_key="k", delay=0.0, manual=True)

        self.assertEqual(calls[1], (["Buddy"], {"Buddy": 2026}))
        self.assertEqual(self._output_years(), {"Buddy": 2026, "Resident Evil": 2002})
        ws = openpyxl.load_workbook(self.output).active
        hm = get_header_map(ws)
        self.assertEqual(ws.cell(2, hm["Reviews"]).value, 22)


if __name__ == "__main__":
    unittest.main()
